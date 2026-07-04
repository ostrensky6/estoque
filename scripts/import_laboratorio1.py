"""Importa Laboratorio1.xlsm para o Supabase de forma idempotente.

O script carrega os inputs crus da planilha de custos laboratoriais e preenche
as tabelas usadas pela engine de custeio do Kontrol. Ele nao apaga dados:
atualiza registros encontrados por chave natural e insere apenas os faltantes.

Uso:
  python scripts/import_laboratorio1.py [caminho_xlsm]
"""

from __future__ import annotations

import datetime as dt
import json
import os
import ssl
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSM = Path(r"D:\Dropbox\ATGC\Custos\1-Laboratorio\Laboratorio1.xlsm")
AFFECTED_TABLES = [
    "parametros",
    "analises",
    "etapas",
    "equipamentos",
    "equipamento_analise",
    "tecnicos",
    "overhead",
    "insumos",
    "insumo_analise",
]
OFERTAVEIS = {
    "Eletrof_vir_hem",
    "Eletrof_vir_tec",
    "Illumina_16S_AC",
    "Illumina_Sh",
    "qPCR_F",
    "qPCR_SF",
    "RTqPCR_RNA_virus_H",
    "RTqPCR_RNA_virus_T",
    "Sanger",
}
INSUMO_ALIASES = {
    "QuantiNovaŽ SYBRŽGreen RT-PCR Kit": "QuantiNova® SYBR®Green RT-PCR Kit (2500 reações)",
    "QuantiNova® SYBR®Green RT-PCR Kit": "QuantiNova® SYBR®Green RT-PCR Kit (2500 reações)",
    "dNTP mix 10 mM. Kit c/ 800 uL (4 x 200 uL)": "dNTP mix 10 mM - 1 mL",
}


def load_env() -> None:
    for env_file in [ROOT / ".env.local", ROOT / ".env"]:
        if not env_file.exists():
            continue
        for raw in env_file.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip("'\""))


def s(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def date_value(value: Any) -> str | None:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return dt.datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    return None


def yes(value: Any) -> bool:
    text = s(value)
    return text is not None and text.lower().startswith("s")


def clean_payload(row: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in row.items() if value is not None}


class SupabaseRest:
    def __init__(self, url: str, key: str) -> None:
        self.base = url.rstrip("/") + "/rest/v1"
        self.headers = {
            "apikey": key,
            "authorization": f"Bearer {key}",
            "content-type": "application/json",
            "accept": "application/json",
        }
        # O ambiente Windows local pode nao ter a cadeia corporativa completa.
        self.context = ssl._create_unverified_context()

    def request(
        self,
        method: str,
        path: str,
        *,
        body: Any | None = None,
        prefer: str | None = None,
    ) -> Any:
        headers = dict(self.headers)
        if prefer:
            headers["Prefer"] = prefer
        data = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
        req = urllib.request.Request(
            f"{self.base}/{path}",
            data=data,
            headers=headers,
            method=method,
        )
        try:
            with urllib.request.urlopen(req, context=self.context, timeout=60) as resp:
                text = resp.read().decode("utf-8")
                return json.loads(text) if text else None
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"{method} {path} falhou: HTTP {exc.code} {detail}") from exc

    def get(self, table: str, query: str = "select=*") -> list[dict[str, Any]]:
        data = self.request("GET", f"{table}?{query}")
        return data if isinstance(data, list) else []

    def insert(self, table: str, row: dict[str, Any]) -> dict[str, Any]:
        data = self.request("POST", table, body=clean_payload(row), prefer="return=representation")
        return data[0] if data else {}

    def update_by_id(self, table: str, row_id: Any, row: dict[str, Any]) -> None:
        query = urllib.parse.urlencode({"id": f"eq.{row_id}"})
        self.request("PATCH", f"{table}?{query}", body=clean_payload(row), prefer="return=minimal")

    def patch(self, table: str, filters: dict[str, Any], row: dict[str, Any]) -> None:
        query = urllib.parse.urlencode({key: f"eq.{value}" for key, value in filters.items()})
        self.request("PATCH", f"{table}?{query}", body=clean_payload(row), prefer="return=minimal")

    def upsert(self, table: str, rows: list[dict[str, Any]], conflict: str) -> None:
        if not rows:
            return
        query = urllib.parse.urlencode({"on_conflict": conflict})
        self.request(
            "POST",
            f"{table}?{query}",
            body=[clean_payload(row) for row in rows],
            prefer="resolution=merge-duplicates,return=minimal",
        )


def workbook_rows(wb: Any, sheet: str):
    return wb[sheet].iter_rows(min_row=2, values_only=True)


def extract(path: Path) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(path, data_only=True, keep_vba=True)

    canon: dict[str, str] = {}
    for row in workbook_rows(wb, "Tempo"):
        code = s(row[1])
        if code:
            canon.setdefault(code.lower(), code)

    def code_norm(value: Any) -> str | None:
        code = s(value)
        return canon.get(code.lower(), code) if code else None

    analises = [
        {
            "codigo": code,
            "nome": code,
            "ativo": True,
            "ofertavel": code in OFERTAVEIS,
        }
        for code in sorted(set(canon.values()))
    ]

    etapa_ordem: defaultdict[str, int] = defaultdict(int)
    etapas: list[dict[str, Any]] = []
    for row in workbook_rows(wb, "Tempo"):
        code = code_norm(row[1])
        if not code:
            continue
        etapa_ordem[code] += 1
        nome_etapa = s(row[2]) or "Etapa"
        nome_atividade = s(row[3]) or "Atividade"
        texto = f"{nome_etapa} {nome_atividade}".lower()
        etapas.append(
            {
                "codigo_analise": code,
                "nome_etapa": nome_etapa,
                "nome_atividade": nome_atividade,
                "execucoes_por_dia": num(row[4]),
                "amostras_por_execucao": num(row[5]),
                "tempo_maquina_h": num(row[7]),
                "tempo_bancada_h": num(row[8]),
                "atividade_opcional": yes(row[10]),
                "tipo_limitacao": s(row[11]),
                "dia_inicio": s(row[12]),
                "dia_fim_max": num(row[13]),
                "ordem": etapa_ordem[code],
                "escopo_operacional": "pos_analise" if "bioinform" in texto else "laboratorio",
            }
        )

    equipamentos: list[dict[str, Any]] = []
    for row in workbook_rows(wb, "Equipamentos"):
        nome = s(row[0])
        if not nome:
            continue
        equipamentos.append(
            {
                "nome": nome,
                "quantidade": num(row[1]) or 1,
                "custo_unitario": num(row[2]) or 0,
                "data_aquisicao": date_value(row[3]),
                "possui": yes(row[5]),
                "vida_util_anos": num(row[6]),
                "percentual_manutencao_anual": num(row[7]) or 0,
            }
        )

    equipamento_analise: list[dict[str, Any]] = []
    for row in workbook_rows(wb, "Equipamento_Analise"):
        nome = s(row[0])
        code = code_norm(row[1])
        if nome and code:
            equipamento_analise.append(
                {
                    "equipamento_nome": nome,
                    "codigo_analise": code,
                    "peso_alocacao": num(row[2]) or 0,
                }
            )

    tecnicos: list[dict[str, Any]] = []
    for row in workbook_rows(wb, "Tecnicos"):
        nome = s(row[1])
        if nome:
            tecnicos.append(
                {
                    "processo": s(row[0]),
                    "nome": nome,
                    "valor_mes": num(row[2]) or 0,
                    "horas_mes_base": num(row[3]) or 170,
                    "percentual_dedicado": num(row[5]) or 0,
                }
            )

    overhead: list[dict[str, Any]] = []
    for row in workbook_rows(wb, "Overhead"):
        item = s(row[0])
        if item:
            overhead.append(
                {
                    "item": item,
                    "custo_mensal": num(row[1]) or 0,
                    "percentual_compensada": num(row[2]) or 100,
                    "horas_bancada_mes": num(row[3]) or 450,
                }
            )

    insumos: list[dict[str, Any]] = []
    seen_specs: set[str] = set()
    for row in workbook_rows(wb, "MC"):
        spec = s(row[1])
        if not spec or spec in seen_specs:
            continue
        seen_specs.add(spec)
        insumos.append(
            {
                "nome_item": s(row[0]),
                "especificacao": spec,
                "custo_unitario": num(row[2]),
                "quantidade_embalagem": num(row[4]),
                "unidade": s(row[6]),
                "custo_total_embalagem": num(row[8]),
                "data_aquisicao": date_value(row[10]),
            }
        )

    insumo_analise: list[dict[str, Any]] = []
    for row in workbook_rows(wb, "MCA"):
        code = code_norm(row[1])
        if not code:
            continue
        spec = s(row[5])
        insumo_analise.append(
            {
                "codigo_analise": code,
                "nome_etapa": s(row[2]) or "Etapa",
                "nome_atividade": s(row[3]) or "Atividade",
                "especificacao_insumo": spec,
                "unidade": s(row[7]),
                "grupo_escolha": s(row[8]),
                "quantidade_por_amostra": num(row[9]),
                "modo_cobranca": s(row[11]),
                "_insumo_lookup": INSUMO_ALIASES.get(spec or "", spec),
            }
        )

    parametros = [
        {
            "chave": "dias_uteis_ano",
            "valor": 222,
            "unidade": "dias",
            "descricao": "Dias uteis/ano p/ rateio de equipamentos",
        },
        {
            "chave": "horas_mes_tecnico",
            "valor": 170,
            "unidade": "h",
            "descricao": "Horas-base mensais por tecnico",
        },
        {
            "chave": "horas_bancada_mes",
            "valor": 450,
            "unidade": "h",
            "descricao": "Horas de bancada/mes p/ rateio de overhead",
        },
        {"chave": "margem_lucro", "valor": 0, "unidade": "%", "descricao": "Margem de lucro sobre o custo total"},
        {"chave": "impostos", "valor": 0, "unidade": "%", "descricao": "Impostos sobre a venda"},
        {"chave": "taxas", "valor": 0, "unidade": "%", "descricao": "Taxas administrativas"},
        {"chave": "fundo_reserva", "valor": 0, "unidade": "%", "descricao": "Fundo de reserva"},
        {"chave": "fundo_investimento", "valor": 0, "unidade": "%", "descricao": "Fundo de investimento"},
    ]

    return {
        "parametros": parametros,
        "analises": analises,
        "etapas": etapas,
        "equipamentos": equipamentos,
        "equipamento_analise": equipamento_analise,
        "tecnicos": tecnicos,
        "overhead": overhead,
        "insumos": insumos,
        "insumo_analise": insumo_analise,
    }


def key_tuple(row: dict[str, Any], fields: list[str]) -> tuple[Any, ...]:
    return tuple(row.get(field) for field in fields)


def sync_by_key(
    db: SupabaseRest,
    table: str,
    source: list[dict[str, Any]],
    key_fields: list[str],
    existing: list[dict[str, Any]] | None = None,
) -> dict[str, int]:
    existing_rows = existing if existing is not None else db.get(table, "select=*")
    by_key = {key_tuple(row, key_fields): row for row in existing_rows}
    inserted = 0
    updated = 0
    for row in source:
        key = key_tuple(row, key_fields)
        current = by_key.get(key)
        if current:
            db.update_by_id(table, current["id"], row)
            updated += 1
        else:
            created = db.insert(table, row)
            if created:
                by_key[key] = created
            inserted += 1
    return {"inserted": inserted, "updated": updated}


def sync_insumo_analise(
    db: SupabaseRest,
    source: list[dict[str, Any]],
    insumo_ids: dict[str, int],
    etapa_ids: dict[tuple[str, str, str], int],
) -> dict[str, int]:
    fields = [
        "codigo_analise",
        "nome_etapa",
        "nome_atividade",
        "especificacao_insumo",
        "unidade",
        "grupo_escolha",
        "quantidade_por_amostra",
        "modo_cobranca",
    ]
    existing = db.get("insumo_analise", "select=*")
    buckets: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in existing:
        buckets[key_tuple(row, fields)].append(row)

    used_existing: Counter[tuple[Any, ...]] = Counter()
    inserted = updated = unresolved_insumo = unresolved_etapa = 0
    for original in source:
        lookup = original.pop("_insumo_lookup", None)
        row = dict(original)
        row["insumo_id"] = insumo_ids.get(lookup) if lookup else None
        etapa_key = (row["codigo_analise"], row["nome_etapa"], row["nome_atividade"])
        row["etapa_id"] = etapa_ids.get(etapa_key)
        if row["especificacao_insumo"] and not row["insumo_id"]:
            unresolved_insumo += 1
        if not row["etapa_id"]:
            unresolved_etapa += 1

        key = key_tuple(row, fields)
        candidates = buckets.get(key, [])
        index = used_existing[key]
        if index < len(candidates):
            db.update_by_id("insumo_analise", candidates[index]["id"], row)
            used_existing[key] += 1
            updated += 1
        else:
            created = db.insert("insumo_analise", row)
            if created:
                buckets[key].append(created)
                used_existing[key] += 1
            inserted += 1
    return {
        "inserted": inserted,
        "updated": updated,
        "unresolved_insumo": unresolved_insumo,
        "unresolved_etapa": unresolved_etapa,
    }


def snapshot(db: SupabaseRest, output_dir: Path) -> dict[str, int]:
    output_dir.mkdir(parents=True, exist_ok=True)
    data = {
        "createdAt": dt.datetime.now(dt.timezone.utc).isoformat(),
        "tables": {table: db.get(table, "select=*") for table in AFFECTED_TABLES},
    }
    (output_dir / "pre-import-snapshot.json").write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return {table: len(rows) for table, rows in data["tables"].items()}


def main() -> None:
    load_env()
    extract_json = "--extract-json" in sys.argv
    args = [arg for arg in sys.argv[1:] if arg != "--extract-json"]
    xlsm = Path(args[0]) if args else DEFAULT_XLSM
    if not xlsm.exists():
        raise SystemExit(f"Planilha nao encontrada: {xlsm}")

    if extract_json:
        print(json.dumps(extract(xlsm), ensure_ascii=False))
        return

    url = os.environ.get("NEXT_PUBLIC_SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY") or os.environ.get("NEXT_PUBLIC_SUPABASE_ANON_KEY")
    if not url or not key:
        raise SystemExit("NEXT_PUBLIC_SUPABASE_URL e uma chave Supabase precisam estar no .env.local.")

    db = SupabaseRest(url, key)
    data = extract(xlsm)
    output_dir = ROOT / "output" / "laboratorio-import"
    before = snapshot(db, output_dir)

    summary: dict[str, Any] = {"before": before, "source": {key: len(value) for key, value in data.items()}, "operations": {}}

    db.upsert("parametros", data["parametros"], "chave")

    existing_analises = {row["codigo"]: row for row in db.get("analises", "select=*")}
    analises_to_insert = [row for row in data["analises"] if row["codigo"] not in existing_analises]
    if analises_to_insert:
        db.upsert("analises", analises_to_insert, "codigo")
    summary["operations"]["analises"] = {"inserted": len(analises_to_insert), "preserved_existing": len(data["analises"]) - len(analises_to_insert)}

    db.upsert("equipamentos", data["equipamentos"], "nome")
    db.upsert("insumos", data["insumos"], "especificacao")

    summary["operations"]["etapas"] = sync_by_key(db, "etapas", data["etapas"], ["codigo_analise", "ordem"])
    summary["operations"]["tecnicos"] = sync_by_key(db, "tecnicos", data["tecnicos"], ["nome", "processo"])
    summary["operations"]["overhead"] = sync_by_key(db, "overhead", data["overhead"], ["item"])

    equipamentos = db.get("equipamentos", "select=id,nome")
    equipamento_ids = {row["nome"]: row["id"] for row in equipamentos}
    equip_alloc = [
        {
            "equipamento_id": equipamento_ids[row["equipamento_nome"]],
            "codigo_analise": row["codigo_analise"],
            "peso_alocacao": row["peso_alocacao"],
        }
        for row in data["equipamento_analise"]
        if row["equipamento_nome"] in equipamento_ids
    ]
    db.upsert("equipamento_analise", equip_alloc, "equipamento_id,codigo_analise")

    insumos = db.get("insumos", "select=id,especificacao")
    insumo_ids = {row["especificacao"]: row["id"] for row in insumos}
    etapas = db.get("etapas", "select=id,codigo_analise,nome_etapa,nome_atividade")
    etapa_counts = Counter((row["codigo_analise"], row["nome_etapa"], row["nome_atividade"]) for row in etapas)
    etapa_ids = {
        (row["codigo_analise"], row["nome_etapa"], row["nome_atividade"]): row["id"]
        for row in etapas
        if etapa_counts[(row["codigo_analise"], row["nome_etapa"], row["nome_atividade"])] == 1
    }
    summary["operations"]["insumo_analise"] = sync_insumo_analise(db, data["insumo_analise"], insumo_ids, etapa_ids)

    after = {table: len(db.get(table, "select=*")) for table in AFFECTED_TABLES}
    summary["after"] = after

    pendencias = db.get(
        "v_insumo_analise_pendencias",
        "select=id,codigo_analise,nome_etapa,nome_atividade,especificacao_insumo,status_vinculo",
    )
    summary["pendencias"] = {
        "total": len(pendencias),
        "por_status": dict(Counter(row["status_vinculo"] for row in pendencias)),
        "amostra": pendencias[:20],
    }

    (output_dir / "import-summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
